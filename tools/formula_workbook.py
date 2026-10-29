"""Build the Formula workbook template, `copilot/formula_workbook.xlsx`.

Run it on `dev` after changing the steps below; `python tools/copilot_files.py` runs it too:

    python tools/formula_workbook.py

A test fails while the committed template differs from what this builds, compared by its
sheets, Tables and cells: a zip's times change with every build, so its bytes would too.

A Formula workbook is laid out so Copilot can read its formulas as SQL: a sample of raw data in
Excel Tables, each calculation in a Table's formulas or a Power Query step, and fixed values in
a Table of their own. A person moves a presentation workbook's data and formulas into a copy of
this one. It holds the steps and the two Tables whose headers are fixed, Parameters and Keys,
empty: no data Table, no formula and no row, since a filled-in example is something to
half-edit and forget.

It is the one file main ships that main can't rebuild: openpyxl, which writes it, is pinned in
requirements-dev.txt only, and the export copies the file as dev holds it.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "copilot" / "formula_workbook.xlsx"

# The steps sheet, one paragraph a row. A heading is a line that ends with a colon.
STEPS = (
    "Formula workbook: how to fill it in",
    "Copilot turns this workbook into Hive, and the Hive into Toolbox code, without anyone "
    "explaining what the workbook is for. Make a copy of this file for each piece of work.",
    "Your data:",
    "1. On a sheet of its own, paste a sample of the raw data as values only (Home > Paste > "
    "Paste Values): one header row, one row per record, no blank rows, no merged cells.",
    "2. Click inside it, press Ctrl+T and tick \"My table has headers\". Its name, under Table "
    "Design > Table Name, can be whatever you like: Copilot keeps the names you give, and "
    "marks any it can't match to a Hive table for you to fill in.",
    "Your calculations:",
    "3. Add a column to the Table for each calculation, and write its formula once, in the "
    "first row, by column name, such as =[@units]*[@price]. Excel fills it down the column.",
    "4. A calculation that reads another Table, such as an XLOOKUP or a SUMIFS, is written the "
    "same way, naming the other Table: =XLOOKUP([@product_id], products[product_id], "
    "products[name]).",
    "5. A Power Query step can do a calculation instead of a formula: Data > From Table/Range, "
    "then Close & Load. Copilot reads its steps too.",
    "6. Move each fixed value a formula uses, such as a rate typed in $B$3, into the Parameters "
    "Table, one row a value, and read it by name: =XLOOKUP(\"tax_rate\", Parameters[name], "
    "Parameters[value]).",
    "Values typed over a formula:",
    "7. A value typed over a formula in one row is converted as it is, for that row. Copilot "
    "needs to know which column or columns tell that Table's rows apart: add a row to the Keys "
    "Table, with the Table's name and those columns, separated by commas. A Table with no "
    "value typed over a formula needs no row in Keys.",
    "Before you hand it over:",
    "8. Take out anything you don't want converted: a Table, a column, a value typed over a "
    "formula. Whatever is left in a Table converts.",
    "9. Charts, PivotTables and anything outside a Table are left alone, and Copilot names "
    "them. Titles and notes can stay.",
    "10. Save the workbook, attach it in a Copilot chat with file 11 loaded, and ask Copilot to "
    "convert it.",
)
# The Tables whose headers are fixed, each on a sheet of its own name.
TABLES = {"Parameters": ("name", "value"), "Keys": ("table", "key_columns")}


def template() -> openpyxl.Workbook:
    """The Formula workbook template, as this builds it."""
    workbook = openpyxl.Workbook()
    steps = workbook.active
    steps.title = "Steps"
    steps.column_dimensions["A"].width = 100
    for row, paragraph in enumerate(STEPS, 1):
        cell = steps.cell(row=row, column=1, value=paragraph)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if row == 1 or paragraph.endswith(":"):
            cell.font = Font(bold=True, size=14 if row == 1 else 11)
    for name, headers in TABLES.items():
        sheet = workbook.create_sheet(name)
        for column, header in enumerate(headers, 1):
            sheet.cell(row=1, column=column, value=header)
            sheet.column_dimensions[chr(ord("A") + column - 1)].width = 24
        # Excel needs a row under the headers, left empty here.
        table = Table(displayName=name, ref="A1:B2")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)
    return workbook


def contents(workbook: openpyxl.Workbook) -> dict[str, object]:
    """What a person sees in a workbook: each sheet's cells by place, and each Table's place."""
    return {sheet.title: {"cells": {cell.coordinate: cell.value for row in sheet.iter_rows()
                                    for cell in row if cell.value is not None},
                          "tables": {name: sheet.tables[name].ref for name in sheet.tables}}
            for sheet in workbook.worksheets}


def main() -> int:
    """Write the template into copilot/."""
    template().save(TEMPLATE)
    print(f"Wrote {TEMPLATE.relative_to(ROOT).as_posix()}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
