"""Build the example Formula workbooks, `copilot/examples/*.xlsx`.

Run it on `dev` after changing an example; `python tools/copilot_files.py` runs it too:

    python tools/formula_workbook_examples.py

The empty Formula workbook, `copilot/formula_workbook.xlsx`, stays empty, as one to copy. These
show it filled in, each a copy of it with sample rows from the Example database's tables, and a
last sheet, Hive, giving the conversion Copilot files 11 and 12 ask for: the map of sources,
steps and finals, then each final's Hive, with `<database>` to fill in. A person can run that
Hive on their own warehouse, or on the Example database with `<database>` as ops.

openpyxl writes no Power Query, so an example's queries are on a sheet of their own, Power Query,
as their M, with the steps for adding them in Excel. A test fails while a committed example
differs from what this builds, and runs each example's Hive.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

from excel_eval_tasks import ExcelTask, jobs_table, runs_table, table
from formula_workbook import ExcelTable, FormulaWorkbook, build, contents

ROOT = Path(__file__).resolve().parent.parent
EXAMPLES_FOLDER = ROOT / "copilot" / "examples"
DATABASE = "<database>"

OK_RUNS = ("ok AS (SELECT run_id, job_id, duration_mins FROM ops.job_runs WHERE status = "
           "'SUCCESS')")

# Each example, by its file's name: the workbook, and its finals' Hive, each table in ops, as the
# Example database holds it. `maps` is the map file 11 asks a reply to start with.
EXAMPLES: dict[str, tuple[ExcelTask, str]] = {
    "01_formulas.xlsx": (ExcelTask(
        "E01", "example", (),
        FormulaWorkbook((
            runs_table(("run_id", "job_id", "status", "duration_mins"), formulas={
                "hours": "=[@duration_mins]/60",
                "job_name": "=XLOOKUP([@job_id],jobs[job_id],jobs[job_name])",
                "cost": '=[@duration_mins]*XLOOKUP("rate",Parameters[name],Parameters[value])'},
                typed_over={(103, "hours"): 9}),
            jobs_table(("job_id", "job_name", "team"))),
            parameters=(("rate", 0.5),), keys=(("job_runs", "run_id"),), as_typed=True),
        {"job_runs": "SELECT r.run_id, r.job_id, r.status, r.duration_mins, CASE WHEN r.run_id "
                     "= 103 THEN 9 ELSE r.duration_mins / 60 END AS hours, j.job_name, "
                     "r.duration_mins * 0.5 AS cost\nFROM ops.job_runs r\nLEFT JOIN ops.jobs j "
                     "ON r.job_id = j.job_id"}),
        "Sources: job_runs, jobs. Finals: job_runs (run_id, job_id, status, duration_mins, hours, "
        "job_name, cost). Typed over: hours where run_id is 103, matched on Keys' run_id. "
        "Parameters: rate, 0.5. Fill in: <database>."),
    "02_summary_and_total_row.xlsx": (ExcelTask(
        "E02", "example", (),
        FormulaWorkbook((
            runs_table(("run_id", "job_id", "status", "duration_mins")),
            ExcelTable("by_job", ("job_id", "total_mins", "runs", "failed", "longest"),
                       "SELECT DISTINCT job_id FROM ops.job_runs ORDER BY job_id", {
                           "total_mins": "=SUMIFS(job_runs[duration_mins],job_runs[job_id],"
                                         "[@job_id])",
                           "runs": "=COUNTIFS(job_runs[job_id],[@job_id])",
                           "failed": '=COUNTIFS(job_runs[job_id],[@job_id],job_runs[status],'
                                     '"FAILED")',
                           "longest": "=MAXIFS(job_runs[duration_mins],job_runs[job_id],"
                                      "[@job_id])"},
                       totals={"job_id": "Total",
                               "total_mins": "=SUBTOTAL(109,by_job[total_mins])",
                               "runs": "=SUBTOTAL(109,by_job[runs])",
                               "failed": "=SUBTOTAL(109,by_job[failed])",
                               "longest": "=SUBTOTAL(104,by_job[longest])"})), as_typed=True),
        {"by_job": "SELECT job_id, SUM(duration_mins) AS total_mins, COUNT(*) AS runs,\n"
                   "  COUNT(CASE WHEN status = 'FAILED' THEN 1 END) AS failed, "
                   "MAX(duration_mins) AS longest\nFROM ops.job_runs\nGROUP BY job_id WITH "
                   "ROLLUP"}),
        "Sources: job_runs. Finals: by_job (job_id, total_mins, runs, failed, longest), a "
        "summary: a GROUP BY of job_runs. Its Total Row adds up sums and counts and takes the "
        "MAX of the MAXes, so it is the Subtotal row WITH ROLLUP adds, its job_id NULL. Fill "
        "in: <database>."),
    "03_power_query.xlsx": (ExcelTask(
        "E03", "example", (),
        FormulaWorkbook(
            (runs_table(),
             ExcelTable("ok_runs", ("run_id", "job_id", "duration_mins", "hours"),
                        "SELECT run_id, job_id, duration_mins FROM ops.job_runs WHERE status = "
                        "'SUCCESS' ORDER BY run_id", {"hours": "=[@duration_mins]/60"},
                        loaded_by="ok_runs"),
             ExcelTable("ok_by_job", ("job_id", "total_mins"),
                        "SELECT job_id, SUM(duration_mins) FROM ops.job_runs WHERE status = "
                        "'SUCCESS' GROUP BY job_id ORDER BY job_id", loaded_by="ok_by_job")),
            queries=(("ok_runs", f"""let
    Source = {table("job_runs")},
    Ok = Table.SelectRows(Source, each [status] = "SUCCESS"),
    Kept = Table.SelectColumns(Ok, {{"run_id", "job_id", "duration_mins"}})
in
    Kept"""), ("ok_by_job", """let
    Grouped = Table.Group(ok_runs, {"job_id"}, {
        {"total_mins", each List.Sum([duration_mins]), type number}})
in
    Grouped""")), as_typed=True),
        {"ok_runs": f"WITH {OK_RUNS}\nSELECT run_id, job_id, duration_mins, duration_mins / 60 "
                    "AS hours\nFROM ok",
         "ok_by_job": f"WITH {OK_RUNS}\nSELECT job_id, SUM(duration_mins) AS total_mins\nFROM ok"
                      "\nGROUP BY job_id"}),
        "Sources: job_runs. Steps: query ok_runs, as query ok_by_job reads it, so it is a CTE "
        "in each final's query. Finals: Table ok_runs (run_id, job_id, duration_mins, hours): "
        "query ok_by_job reads the query, not the Table, and no formula reads the Table, so it "
        "is a final, its query's rows and its hours column; Table ok_by_job (job_id, "
        "total_mins). Fill in: <database>."),
}

POWER_QUERY_STEPS = (
    "This example's Power Queries, which the tool that made this file can't save in it, so the "
    "Tables ok_runs and ok_by_job hold the rows the queries give, typed in. To see the workbook "
    "as it should be, add each query in Excel, in order: Data > Get Data > From Other Sources > "
    "Blank Query, Home > Advanced Editor, paste its M below, and name it as below. On the sheet "
    "of its name, delete the typed-in Table (select it and press Delete, then delete the empty "
    "Table), then Home > Close & Load To > Table, at A1 of that sheet. For ok_runs, type the "
    "formula column again beside the loaded rows: hours, =[@duration_mins]/60. Excel keeps it "
    "when the query refreshes.")


def shown(sql: str) -> str:
    """A final's Hive as the example shows it: its tables in `<database>`, not ops."""
    return sql.replace("ops.", f"{DATABASE}.")


def example_workbook(task: ExcelTask, map_text: str) -> openpyxl.Workbook:
    """An example: the filled-in Formula workbook, its Power Queries if it has any, and the Hive
    files 11 and 12 turn it into."""
    workbook = build(task.workbook)
    if task.workbook.queries:
        _sheet_of_code(workbook, "Power Query", [POWER_QUERY_STEPS],
                       [(f"Query {name}:", m) for name, m in task.workbook.queries])
    hive = _sheet_of_code(workbook, "Hive", [
        "What Copilot gives with files 11, 12, 03, 04 and 05 loaded: the map, then a query per "
        "final. Fill in <database> with your database's name. Delete this sheet, and the Power "
        "Query sheet if there is one, before you try the workbook on Copilot yourself.",
        map_text],
        [(f"{name}:", shown(sql)) for name, sql in task.finals.items()])
    hive.insert_rows(1)
    hive["A1"] = "The Hive this workbook converts to"
    hive["A1"].font = Font(bold=True, size=14)
    return workbook


# A sheet of text and code: its one column's width, in characters.
WIDTH = 100


def _sheet_of_code(workbook: openpyxl.Workbook, title: str, paragraphs: list[str],
                   blocks: list[tuple[str, str]]) -> openpyxl.worksheet.worksheet.Worksheet:
    """A sheet of one wide column: the paragraphs, a blank row after each, then each block's
    label in bold over its code in a fixed-width font, a blank row after each."""
    sheet = workbook.create_sheet(title)
    sheet.column_dimensions["A"].width = WIDTH
    wrap = Alignment(wrap_text=True, vertical="top")
    row = 1
    for paragraph in paragraphs:
        sheet.cell(row=row, column=1, value=paragraph).alignment = wrap
        row += 2
    for label, code in blocks:
        sheet.cell(row=row, column=1, value=label).font = Font(bold=True)
        cell = sheet.cell(row=row + 1, column=1, value=code)
        cell.alignment, cell.font = wrap, Font(name="Consolas")
        row += 3
    return sheet


def main() -> int:
    """Write each example into copilot/examples/, unless it already holds what this builds."""
    EXAMPLES_FOLDER.mkdir(exist_ok=True)
    for name, (task, map_text) in EXAMPLES.items():
        path = EXAMPLES_FOLDER / name
        built = example_workbook(task, map_text)
        if path.exists() and contents(openpyxl.load_workbook(path)) == contents(built):
            continue
        built.save(path)
        print(f"Wrote {path.relative_to(ROOT).as_posix()}.")
    print(f"{len(EXAMPLES)} example workbooks in copilot/examples/.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
