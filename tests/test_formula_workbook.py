"""The empty Formula workbook, `copilot/formula_workbook.xlsx` (Toolbox 7.2, ticket 15), and the
examples of it filled in, `copilot/examples/`, each with the Hive it converts to.

main can't rebuild it: tools/formula_workbook.py builds it with openpyxl, which only dev
installs, and the export copies the file. It holds a sheet of steps and two empty Excel Tables,
Parameters and Keys, and nothing a person could half-edit: no data Table, no formula, no row.
The offline policy and the code audit read it as they read the repo (tests/repo/).
"""

from __future__ import annotations

import openpyxl
import pytest

import excel_eval
import excel_eval_tasks
import formula_workbook
import formula_workbook_examples
from test_copilot_files import AVOIDED

EMPTY = formula_workbook.EMPTY
# The glossary's words to avoid for a Formula workbook.
NOT_ITS_NAME = ("state table", "Excel template", "mapping sheet")


@pytest.fixture(scope="module")
def committed() -> openpyxl.Workbook:
    return openpyxl.load_workbook(EMPTY)


def test_it_is_what_the_tool_builds(committed: openpyxl.Workbook) -> None:
    # Compared by what a person sees, not by bytes: a zip's times change with every build.
    assert formula_workbook.contents(committed) == formula_workbook.contents(
        formula_workbook.empty_workbook())


def test_it_holds_the_steps_and_two_empty_tables(committed: openpyxl.Workbook) -> None:
    assert committed.sheetnames == ["Steps", "Parameters", "Keys"]
    tables = {name: (sheet.title, sheet.tables[name].ref) for sheet in committed.worksheets
              for name in sheet.tables}
    assert tables == {"Parameters": ("Parameters", "A1:B2"), "Keys": ("Keys", "A1:B2")}
    headers = {sheet.title: [sheet["A1"].value, sheet["B1"].value]
               for sheet in committed.worksheets[1:]}
    assert headers == {"Parameters": ["name", "value"], "Keys": ["table", "key_columns"]}
    for sheet in committed.worksheets[1:]:
        assert [cell.value for row in sheet.iter_rows(min_row=2) for cell in row
                if cell.value is not None] == [], sheet.title


def test_each_tables_sheet_says_what_it_is_for_beside_it(committed: openpyxl.Workbook) -> None:
    for sheet in committed.worksheets[1:]:
        note = sheet["D1"].value
        assert note.startswith(f"{sheet.title}: "), sheet.title
        assert [cell.coordinate for row in sheet.iter_rows() for cell in row
                if cell.value is not None] == ["A1", "B1", "D1"], sheet.title
    steps = [cell.value for row in committed["Steps"].iter_rows() for cell in row]
    about = steps.index("The Parameters and Keys sheets:")
    assert about < steps.index("Your data:")
    assert "Parameters holds" in steps[about + 1] and "Keys is" in steps[about + 1]


def test_no_cell_holds_a_formula(committed: openpyxl.Workbook) -> None:
    formulas = [f"{sheet.title}!{cell.coordinate}" for sheet in committed.worksheets
                for row in sheet.iter_rows() for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert formulas == []


def test_the_steps_use_no_word_the_glossary_avoids(committed: openpyxl.Workbook) -> None:
    text = "\n".join(str(cell.value) for row in committed["Steps"].iter_rows() for cell in row
                     if cell.value is not None)
    assert [word for word in (*AVOIDED, *NOT_ITS_NAME) if word.lower() in text.lower()] == []


@pytest.mark.parametrize("name", list(formula_workbook_examples.EXAMPLES))
def test_each_example_is_what_the_tool_builds_and_its_hive_runs(name: str) -> None:
    task, map_text = formula_workbook_examples.EXAMPLES[name]
    committed = openpyxl.load_workbook(formula_workbook_examples.EXAMPLES_FOLDER / name)
    assert formula_workbook.contents(committed) == formula_workbook.contents(
        formula_workbook_examples.example_workbook(task, map_text))
    assert committed.sheetnames[:3] == ["Steps", "Parameters", "Keys"]
    assert committed.sheetnames[-1] == "Hive"
    assert ("Power Query" in committed.sheetnames) == bool(task.workbook.queries)
    shown = [cell.value for row in committed["Hive"].iter_rows() for cell in row]
    for sql in task.finals.values():
        assert formula_workbook_examples.shown(sql) in shown
        assert "ops." not in formula_workbook_examples.shown(sql)
    rows = excel_eval.reference_rows(excel_eval.stand_in_hive(), task)
    assert all(rows.values()), rows


def _what_excel_would_repair(workbook: openpyxl.Workbook) -> list[str]:
    """Each place a Table's definition disagrees with its cells, which Excel "repairs" on
    opening by dropping the Table: its columns must be its headers, its filter must stop above
    its Total Row, and each Total Row cell must be what its column says, a label or a
    function."""
    found = []
    for sheet in workbook.worksheets:
        for table in sheet.tables.values():
            rows = sheet[table.ref]
            where = f"{sheet.title}!{table.displayName}"
            if ([column.name for column in table.tableColumns]
                    != [cell.value for cell in rows[0]]):
                found.append(f"{where}: its columns aren't its headers")
            totals = rows[-1] if table.totalsRowCount else ()
            data = rows[:-1] if totals else rows
            if table.autoFilter.ref != f"{data[0][0].coordinate}:{data[-1][-1].coordinate}":
                found.append(f"{where}: its filter is {table.autoFilter.ref}")
            for column, cell in zip(table.tableColumns, totals):
                function = column.totalsRowFunction
                held = (f"={column.totalsRowFormula.attr_text}" if function == "custom" else
                        f"=SUBTOTAL({formula_workbook.SUBTOTAL_CODES[function]},"
                        f"{table.displayName}[{column.name}])" if function else
                        column.totalsRowLabel)
                if cell.value != held:
                    found.append(f"{where}: {cell.coordinate} is {cell.value!r}, its column "
                                 f"says {held!r}")
    return found


@pytest.mark.parametrize("name", list(formula_workbook_examples.EXAMPLES))
def test_excel_opens_each_example_without_a_repair(name: str) -> None:
    assert _what_excel_would_repair(
        openpyxl.load_workbook(formula_workbook_examples.EXAMPLES_FOLDER / name)) == []


def test_excel_would_open_each_eval_workbook_without_a_repair() -> None:
    found = {task.id: _what_excel_would_repair(formula_workbook.build(task.workbook))
             for task in excel_eval_tasks.TASKS}
    assert {task: problems for task, problems in found.items() if problems} == {}


def test_a_total_row_cell_that_is_not_a_subtotal_of_its_column_is_custom() -> None:
    built = formula_workbook.build(formula_workbook.FormulaWorkbook((
        formula_workbook.ExcelTable("t", ("a", "b"), "SELECT 1 AS a", {"b": "=[@a]*2"},
                                    totals={"a": "=SUBTOTAL(109,t[b])", "b": "=MAX(t[b])"}),)))
    columns = built["t"].tables["t"].tableColumns
    assert [(c.totalsRowFunction, c.totalsRowFormula.attr_text) for c in columns] == [
        ("custom", "SUBTOTAL(109,t[b])"), ("custom", "MAX(t[b])")]
    assert _what_excel_would_repair(built) == []


def test_the_examples_start_from_the_empty_workbook(committed: openpyxl.Workbook) -> None:
    for name in formula_workbook_examples.EXAMPLES:
        example = openpyxl.load_workbook(formula_workbook_examples.EXAMPLES_FOLDER / name)
        assert [cell.value for row in example["Steps"].iter_rows() for cell in row] == [
            cell.value for row in committed["Steps"].iter_rows() for cell in row], name
        for title in ("Parameters", "Keys"):
            assert example[title]["D1"].value == committed[title]["D1"].value, name
